
"""
NutriCare AI regression tests.

Run from the project root:

    python -m pytest -q
"""

from __future__ import annotations

import io

from pathlib import Path

import pytest

from fastapi.testclient import TestClient

from docx import Document

from openpyxl import Workbook

from reportlab.pdfgen.canvas import Canvas


from backend.main import app

from backend.ingestion import parse_file

from backend.retrieval import lexical_retrieve

from backend.evaluation import (
    precision_recall_at_k,
    patient_leakage,
    evaluate_golden_case,
)


# ==========================================
# TEST CONFIGURATION
# ==========================================

CLIENT = TestClient(app)

ROOT = Path(__file__).resolve().parents[1]

CSV = (
    ROOT / "public" / "demo-patients.csv"
).read_bytes()


# ==========================================
# UPLOAD HELPER
# ==========================================

def submit(
    contents: bytes,
    filename: str,
    mode: str = "group",
    confirmed: bool = True,
):

    return CLIENT.post(

        "/api/upload",

        data={

            "mode": mode,

            "demo_confirmed": str(confirmed).lower(),

        },

        files={

            "file": (
                filename,
                contents
            )

        },

    )


# ==========================================
# GROUP CSV UPLOAD
# ==========================================

def test_group_csv_upload_and_patient_isolation():

    response = submit(
        CSV,
        "demo-patients.csv"
    )

    assert response.status_code == 200, response.text

    patients = response.json()["patients"]

    assert len(patients) == 3

    assert patients[0]["id"] == "DEMO101"

    assert patients[1]["id"] == "DEMO102"

    assert all(

        "Sam Taylor" not in chunk["text"]

        for chunk in patients[0]["chunks"]

    )

    assert patients[0]["evidence"]["missing"] is not None


# ==========================================
# SINGLE PATIENT MODE
# ==========================================

def test_single_mode_rejects_multiple_patients():

    result = submit(

        CSV,

        "demo-patients.csv",

        mode="single"

    )

    assert result.status_code == 422


# ==========================================
# PATIENT IDENTIFIER REQUIREMENT
# ==========================================

def test_group_requires_verified_id_column():

    sample = (

        b"patient_name,age,diet\n"

        b"Alex Rivera,45,rice\n"

        b"Sam Taylor,52,beans\n"

    )


    response = submit(

        sample,

        "sample.csv"

    )


    assert response.status_code == 422

    assert "patient_id" in response.json()["detail"]


# ==========================================
# SYNTHETIC DATA CONFIRMATION
# ==========================================

def test_synthetic_attestation_required():

    result = submit(

        CSV,

        "demo-patients.csv",

        confirmed=False

    )


    assert result.status_code == 400


# ==========================================
# FILE SIZE AND SIGNATURE
# ==========================================

def test_upload_size_and_signature():

    large_file = b"X" * (
        3 * 1024 * 1024 + 1
    )


    assert submit(

        large_file,

        "demo.csv"

    ).status_code == 413


    assert submit(

        b"not a pdf",

        "record.pdf",

        "single"

    ).status_code == 400


# ==========================================
# WORD DOCUMENT UPLOAD
# ==========================================

def test_docx_single():

    doc = Document()


    doc.add_paragraph(
        "PATIENT NAME: Taylor Example"
    )

    doc.add_paragraph(
        "AGE: 40"
    )

    doc.add_paragraph(
        "MEDICAL HISTORY: Synthetic study case"
    )

    doc.add_paragraph(
        "DIET HISTORY: Beans and vegetables"
    )


    raw = io.BytesIO()

    doc.save(raw)


    response = submit(

        raw.getvalue(),

        "demo.docx",

        "single"

    )


    assert response.status_code == 200, response.text


    assert (

        response.json()["patients"][0]["name"]

        == "Taylor Example"

    )


# ==========================================
# EXCEL GROUP UPLOAD
# ==========================================

def test_excel_group():

    book = Workbook()

    sheet = book.active


    sheet.append([

        "patient_id",

        "patient_name",

        "age",

        "diet",

    ])


    sheet.append([

        "AA001",

        "Synthetic A",

        23,

        "beans",

    ])


    sheet.append([

        "BB002",

        "Synthetic B",

        25,

        "oats",

    ])


    stream = io.BytesIO()

    book.save(stream)


    response = submit(

        stream.getvalue(),

        "demo.xlsx"

    )


    assert response.status_code == 200, response.text


    assert len(
        response.json()["patients"]
    ) == 2


# ==========================================
# PDF UPLOAD
# ==========================================

def test_pdf_single():

    stream = io.BytesIO()

    pdf = Canvas(stream)


    pdf.drawString(
        40,
        790,
        "PATIENT NAME: Synthetic Sample"
    )


    pdf.drawString(
        40,
        760,
        "AGE: 35"
    )


    pdf.drawString(
        40,
        730,
        "DIET HISTORY: oats and beans"
    )


    pdf.save()


    response = submit(

        stream.getvalue(),

        "demo.pdf",

        "single"

    )


    assert response.status_code == 200, response.text


    assert (

        response.json()["patients"][0]["name"]

        == "Synthetic Sample"

    )


# ==========================================
# RETRIEVAL EVALUATION
# ==========================================

def test_retrieval_returns_matching_patient_chunk():

    chunks = [

        {

            "id": "S1",

            "text": "allergies are not reported",

            "source": "a",

            "location": "1",

        },

        {

            "id": "S2",

            "text": "diet is rice and beans",

            "source": "a",

            "location": "2",

        },

    ]


    assert lexical_retrieve(

        "allergies",

        chunks

    )[0]["id"] == "S1"


    scores = precision_recall_at_k(

        ["S1", "S2"],

        {"S1"},

        k=2

    )


    assert scores["precision_at_k"] == 0.5

    assert scores["recall_at_k"] == 1.0


    assert patient_leakage(

        [
            {
                "patient_id": "B"
            }
        ],

        {"A"}

    ) is True


# ==========================================
# AI CONFIGURATION
# ==========================================

def test_chat_rejects_missing_config_or_invalid_data(
    monkeypatch
):

    monkeypatch.delenv(

        "CLOUDFLARE_API_TOKEN",

        raising=False

    )


    monkeypatch.delenv(

        "CLOUDFLARE_ACCOUNT_ID",

        raising=False

    )


    p = submit(

        CSV,

        "demo-patients.csv"

    ).json()["patients"][0]


    response = CLIENT.post(

        "/api/assess",

        json={

            "patient": p,

            "demo_confirmed": True,

        }

    )


    assert response.status_code == 503


# ==========================================
# PDF REPORT EXPORT
# ==========================================

def test_export_requires_approval_and_returns_pdf():

    p = submit(

        CSV,

        "demo-patients.csv"

    ).json()["patients"][0]


    data = {

        "patient": p,

        "plan": (
            "Synthetic nutrition discussion draft "
            "for fictional clinical review."
        ),

        "reviewer": "Demo Reviewer",

        "demo_confirmed": True,

    }


    response = CLIENT.post(

        "/api/report",

        json={

            **data,

            "approved": False,

        }

    )


    assert response.status_code == 422


    response = CLIENT.post(

        "/api/report",

        json={

            **data,

            "approved": True,

        }

    )


    assert response.status_code == 200


    assert (

        response.headers["content-type"]

        == "application/pdf"

    )


    assert response.content[:5] == b"%PDF-"


    assert (

        "Recommendation.pdf"

        in response.headers["content-disposition"]

    )


# ==========================================
# API HEALTH
# ==========================================

def test_api_works_without_credentials(
    monkeypatch
):

    monkeypatch.delenv(

        "CLOUDFLARE_API_TOKEN",

        raising=False

    )


    monkeypatch.delenv(

        "CLOUDFLARE_ACCOUNT_ID",

        raising=False

    )


    assert CLIENT.get("/").status_code == 200


    status = CLIENT.get(
        "/api/health"
    ).json()


    assert status["status"] == "ok"

    assert status["ai_configured"] is False


    assert CLIENT.get(
        "/demo-patients.csv"
    ).status_code == 200


# ==========================================
# RAG AND LLM EVALUATION
# ==========================================

def test_annotated_rag_llm_evaluation():

    values = evaluate_golden_case(

        retrieved_ids=[
            "S1",
            "S2"
        ],

        relevant_ids={
            "S1"
        },

        cited_ids=[
            "S1"
        ],

        supported_ids={
            "S1"
        },

        answer_relevant=True,

        claim_grounded_labels=[
            True,
            False
        ],

        k=2,

    )


    assert values["context_relevance"] == 0.5

    assert values["answer_relevance"] == 1.0

    assert values["groundedness"] == 0.5

    assert values["citation_accuracy"] == 1.0


# ==========================================
# MULTILINE PDF TABLE EXTRACTION
# ==========================================

def test_multiline_pdf_table_extracts_patient_facts():

    """
    Tests a PDF table that extracts as:

        Label
        Value

    Instead of:

        KEY: VALUE
    """


    data = io.BytesIO()


    pdf = Canvas(

        data,

        pagesize=(595, 842)

    )


    lines = [

        "01 | Patient identification and intake",

        "Full name",

        "Amina Bello (fictional)",

        "Patient ID",

        "DEMO-NUTRI-0001 (synthetic)",

        "Date of birth / age",

        "18 February 1980 / 46 years on the assessment date",

        "Sex recorded for this case",

        "Female",

        "02 | Medical and clinical history",

        "Type 2 diabetes mellitus",

        "Documented diagnosis in 2022.",

        "Food allergy",

        "Patient reports a peanut allergy.",

    ]


    y = 800


    for line in lines:

        pdf.drawString(

            35,

            y,

            line

        )


        y -= 25


    pdf.showPage()


    page_two = [

        "03 | Anthropometrics and vital signs",

        "Height",

        "166 cm",

        "Weight",

        "86.0 kg",

        "Metformin extended release",

        "500 mg twice daily",

        "04 | Laboratory results",

        "Fasting plasma glucose",

        "162 mg/dL",

        "Prior dated results for comparison",

        "HbA1c",

        "7.6 %",

    ]


    for index, line in enumerate(page_two):

        pdf.drawString(

            35,

            800 - 35 * index,

            line

        )


    pdf.save()


    response = submit(

        data.getvalue(),

        "record.pdf",

        "single"

    )


    assert response.status_code == 200, response.text


    person = response.json()["patients"][0]


    assert person["name"] == "Amina Bello"


    assert person["id"] == "DEMO-NUTRI-0001"


    assert "46 years" in person["fields"]["age"]


    assert person["fields"]["weight"] == "86.0 kg"


    assert (

        "peanut allergy"

        in person["fields"]["allergies"]

    )


# ==========================================
# FRONTEND REGRESSION TESTS
# ==========================================

def test_frontend_updated_navigation_and_scroll_design():

    markup = (

        ROOT / "public" / "index.html"

    ).read_text(
        encoding="utf-8"
    )


    css = (

        ROOT / "public" / "style.css"

    ).read_text(
        encoding="utf-8"
    )


    assert 'id="sidebarToggle"' in markup


    assert 'data-view="knowledge"' not in markup


    assert "Open source demo" not in markup


    assert "privacy-tile" not in markup


    assert 'rel="icon"' in markup


    assert "overflow-y: auto" in css


    assert ".sidebar-collapsed" in css