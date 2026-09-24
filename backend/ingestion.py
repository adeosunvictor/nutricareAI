
"""
Document ingestion for NutriCare AI.

Supported formats:
    PDF
    DOCX
    CSV
    XLSX

Responsibilities:
    - Extract document text.
    - Parse structured patient information.
    - Separate individual and group patient records.
    - Preserve source locations.
    - Prepare document chunks for retrieval.

The MVP processes synthetic patient records only.
"""

from __future__ import annotations

import csv
import io
import re
import uuid

from pathlib import Path

from fastapi import HTTPException

from .security import MAX_PATIENTS, normalize_text

from .extraction import extract_document_fields


# ==========================================
# SUPPORTED FIELD ALIASES
# ==========================================

ALIASES = {

    "patient_id": (
        "patient_id",
        "patient id",
        "record_id",
        "record id",
        "mrn",
        "id",
    ),

    "name": (
        "patient_name",
        "patient name",
        "name",
        "full_name",
        "full name",
    ),

    "age": (
        "age",
        "age_years",
    ),

    "sex": (
        "sex",
        "gender",
    ),

    "height": (
        "height",
        "height_cm",
        "height (cm)",
    ),

    "weight": (
        "weight",
        "weight_kg",
        "weight (kg)",
    ),

    "medical_history": (
        "medical_history",
        "medical history",
        "conditions",
        "diagnosis",
        "illness_history",
    ),

    "medications": (
        "medications",
        "medication_history",
        "medication",
        "current_medications",
    ),

    "labs": (
        "labs",
        "lab_results",
        "lab results",
        "laboratory_results",
        "results",
    ),

    "diet": (
        "diet",
        "diet_history",
        "diet history",
        "food_preferences",
        "dietary_history",
    ),

    "allergies": (
        "allergies",
        "food_allergies",
        "food allergies",
    ),

    "budget": (
        "budget",
        "food_budget",
        "financial_constraints",
    ),

}


# ==========================================
# NORMALIZE FIELD NAMES
# ==========================================

def _field_key(value: str) -> str:

    return re.sub(
        r"\s+",
        "_",
        str(value).strip().lower()
    )


ALIAS_MAP = {

    _field_key(alias): field

    for field, aliases in ALIASES.items()

    for alias in aliases

}


# ==========================================
# CREATE PATIENT OBJECT
# ==========================================

def _create_patient(
    name: str,
    patient_id: str,
    source: str
) -> dict:

    return {

        "id": patient_id,

        "name": name,

        "fields": {},

        "chunks": [],

        "source": source,

        "verification": "unverified",

        "plan": "",

        "approved": False,

    }


# ==========================================
# CREATE SOURCE CHUNK
# ==========================================

def _append_chunk(
    patient: dict,
    text: str,
    source: str,
    location: str
) -> None:

    text = normalize_text(
        text,
        1300
    )

    if not text:
        return

    patient["chunks"].append({

        "id": f"S{len(patient['chunks']) + 1}",

        "text": text,

        "source": source,

        "location": location,

    })


# ==========================================
# CONVERT SPREADSHEET ROWS TO PATIENTS
# ==========================================

def _rows_to_patients(
    rows: list[dict],
    source: str,
    mode: str
) -> list[dict]:

    if not rows:

        raise HTTPException(
            422,
            "No data rows were found in the spreadsheet."
        )

    grouped = {}


    for index, row in enumerate(
        rows[:2000],
        2
    ):

        fields = {

            ALIAS_MAP.get(
                _field_key(k),
                _field_key(k)
            ): normalize_text(v, 1300)

            for k, v in row.items()

            if (
                k is not None
                and v is not None
                and str(v).strip()
            )

        }


        patient_id = fields.get(
            "patient_id",
            ""
        ).strip()


        name = fields.get(
            "name",
            ""
        ).strip()


        # Group uploads require patient identifiers.

        if mode == "group" and not patient_id:

            raise HTTPException(
                422,
                f"Row {index} has no patient_id. "
                "Group files must include a unique "
                "patient_id column."
            )


        if not name:

            raise HTTPException(
                422,
                f"Row {index} is missing "
                "patient_name / name."
            )


        key = patient_id or "SINGLE"


        # Create patient if not already encountered.

        if key not in grouped:

            if len(grouped) >= MAX_PATIENTS:

                raise HTTPException(
                    422,
                    f"Demo supports a maximum of "
                    f"{MAX_PATIENTS} patients per upload."
                )


            safe_id = re.sub(
                r"[^A-Za-z0-9_-]",
                "",
                patient_id
            )[:45]


            grouped[key] = _create_patient(

                name,

                safe_id or uuid.uuid4().hex[:12],

                source

            )


        patient = grouped[key]


        # Prevent inconsistent patient identity.

        if (
            patient["name"].casefold()
            != name.casefold()
        ):

            raise HTTPException(
                422,
                f"Patient ID {key} has conflicting names. "
                "Correct the source file."
            )


        # Merge structured fields.

        for field, value in fields.items():

            if field in {
                "name",
                "patient_id"
            }:

                continue


            previous = patient["fields"].get(
                field,
                ""
            )


            if value not in previous.split(" | "):

                patient["fields"][field] = " | ".join(

                    filter(
                        None,
                        (previous, value)
                    )

                )[:2500]


        # Preserve row information for retrieval.

        summary = "; ".join(

            f"{k}: {v}"

            for k, v in fields.items()

            if k not in {
                "name",
                "patient_id"
            }

        )


        _append_chunk(

            patient,

            summary,

            source,

            f"row {index}"

        )


    if (
        mode == "single"
        and len(grouped) != 1
    ):

        raise HTTPException(
            422,
            "Single patient mode requires exactly "
            "one patient. Select group upload."
        )


    return list(grouped.values())


# ==========================================
# CSV PARSER
# ==========================================

def _parse_csv(
    contents: bytes,
    source: str,
    mode: str
) -> list[dict]:

    try:

        content = contents.decode(
            "utf-8-sig"
        )


        reader = csv.DictReader(
            io.StringIO(content)
        )


        if not reader.fieldnames:

            raise ValueError(
                "CSV needs a header row"
            )


        return _rows_to_patients(

            list(reader),

            source,

            mode

        )


    except UnicodeDecodeError as exc:

        raise HTTPException(
            422,
            "CSV must be UTF-8 encoded."
        ) from exc


    except (csv.Error, ValueError) as exc:

        raise HTTPException(
            422,
            f"CSV could not be read: {exc}"
        ) from exc


# ==========================================
# EXCEL PARSER
# ==========================================

def _parse_xlsx(
    contents: bytes,
    source: str,
    mode: str
) -> list[dict]:

    from openpyxl import load_workbook


    try:

        book = load_workbook(

            io.BytesIO(contents),

            read_only=True,

            data_only=True

        )


        all_rows = []


        for sheet in book.worksheets:

            iterator = sheet.iter_rows(
                values_only=True
            )


            headings = next(
                iterator,
                None
            )


            if not headings:
                continue


            for row in iterator:

                if any(

                    value is not None
                    and str(value).strip()

                    for value in row

                ):

                    all_rows.append({

                        str(
                            headings[i] or f"column_{i}"
                        ): value

                        for i, value in enumerate(row)

                        if i < len(headings)

                    })


                if len(all_rows) > 2000:
                    break


        return _rows_to_patients(

            all_rows,

            source,

            mode

        )


    except HTTPException:
        raise


    except Exception as exc:

        raise HTTPException(
            422,
            "Excel file could not be read. "
            "Check that it is a valid XLSX file."
        ) from exc


# ==========================================
# PDF TEXT EXTRACTION
# ==========================================

def _extract_pdf(
    contents: bytes
) -> list[tuple[str, str]]:

    from pypdf import PdfReader


    try:

        reader = PdfReader(
            io.BytesIO(contents)
        )


        if len(reader.pages) > 40:

            raise HTTPException(
                422,
                "Demo supports PDF files of up to 40 pages."
            )


        return [

            (
                page.extract_text() or "",
                f"page {i + 1}"
            )

            for i, page in enumerate(reader.pages)

        ]


    except HTTPException:
        raise


    except Exception as exc:

        raise HTTPException(
            422,
            "Could not extract this PDF. "
            "Scanned-only PDFs are not supported "
            "in the MVP."
        ) from exc


# ==========================================
# WORD DOCUMENT EXTRACTION
# ==========================================

def _extract_docx(
    contents: bytes
) -> list[tuple[str, str]]:

    from docx import Document


    try:

        doc = Document(
            io.BytesIO(contents)
        )


        parts = [

            (
                p.text.strip(),
                f"paragraph {i + 1}"
            )

            for i, p in enumerate(doc.paragraphs)

            if p.text.strip()

        ]


        for i, table in enumerate(doc.tables):

            for j, row in enumerate(table.rows):

                text = "; ".join(

                    cell.text.strip()

                    for cell in row.cells

                    if cell.text.strip()

                )


                if text:

                    parts.append((

                        text,

                        f"table {i + 1}, row {j + 1}"

                    ))


        return parts


    except Exception as exc:

        raise HTTPException(
            422,
            "Could not read this Word document."
        ) from exc


# ==========================================
# DOCUMENT FIELD PATTERNS
# ==========================================

_GROUP_MARKER = re.compile(

    r"(?im)^\s*PATIENT\s+ID\s*:\s*([\w-]{1,45})\s*$"

)


_NAME_MARKER = re.compile(

    r"(?im)^\s*(?:PATIENT\s+NAME|NAME)\s*:\s*([^\n]{1,110})"

)


_SINGLE_FIELD = re.compile(

    r"(?im)^\s*(?:PATIENT\s+NAME|NAME|AGE|SEX|GENDER|HEIGHT|WEIGHT|MEDICAL\s+HISTORY|MEDICATIONS?|LAB(?:ORATORY)?\s+RESULTS?|DIET(?:ARY)?\s+HISTORY|ALLERGIES|BUDGET)\s*:\s*(.+)$"

)


# ==========================================
# PARSE SINGLE OR GROUP DOCUMENT
# ==========================================

def _parse_text(
    parts: list[tuple[str, str]],
    source: str,
    mode: str
) -> list[dict]:

    whole = "\n".join(

        text

        for text, _ in parts

    )


    if len(whole.strip()) < 25:

        raise HTTPException(
            422,
            "The document contains too little "
            "selectable text. Please use a "
            "text-based PDF or DOCX."
        )


    # ======================================
    # GROUP DOCUMENT
    # ======================================

    if mode == "group":

        matches = list(
            _GROUP_MARKER.finditer(whole)
        )


        if not matches:

            raise HTTPException(
                422,
                "Group PDF/DOCX must contain "
                "PATIENT ID: <id> on a separate "
                "line at the start of every "
                "patient section."
            )


        if whole[:matches[0].start()].strip():

            raise HTTPException(
                422,
                "Group document contains text "
                "before its first PATIENT ID section. "
                "Move that text inside a patient section."
            )


        if len(matches) > MAX_PATIENTS:

            raise HTTPException(
                422,
                f"Demo supports no more than "
                f"{MAX_PATIENTS} patients."
            )


        patients = []

        seen_ids = set()


        for i, match in enumerate(matches):

            end = (

                matches[i + 1].start()

                if i + 1 < len(matches)

                else None

            )


            block = whole[
                match.start():end
            ]


            patient_id = match.group(1)


            if patient_id in seen_ids:

                raise HTTPException(
                    422,
                    f"Duplicate patient ID "
                    f"{patient_id} in group document."
                )


            seen_ids.add(patient_id)


            name_match = _NAME_MARKER.search(
                block
            )


            if not name_match:

                raise HTTPException(
                    422,
                    f"Patient {patient_id} needs "
                    "a PATIENT NAME: line."
                )


            patient = _create_patient(

                name_match.group(1).strip(),

                patient_id,

                source

            )


            for match_field in _SINGLE_FIELD.finditer(
                block
            ):

                field_name = match_field.group(0).split(
                    ":",
                    1
                )[0]


                field = ALIAS_MAP.get(

                    _field_key(field_name),

                    _field_key(field_name)

                )


                if field != "name":

                    patient["fields"][field] = normalize_text(

                        match_field.group(1),

                        500

                    )


            fragments = re.findall(

                r".{1,850}(?:\s|$)",

                block,

                re.S

            )


            for j, fragment in enumerate(fragments):

                _append_chunk(

                    patient,

                    fragment,

                    source,

                    f"patient section {i + 1}, part {j + 1}"

                )


            patients.append(patient)


        return patients


    # ======================================
    # SINGLE DOCUMENT
    # ======================================

    if len(_GROUP_MARKER.findall(whole)) > 1:

        raise HTTPException(
            422,
            "This document contains multiple "
            "patient sections. Choose group upload."
        )


    structured = extract_document_fields(
        parts
    )


    name_match = _NAME_MARKER.search(
        whole
    )


    name = (

        structured["name"]

        or (

            name_match.group(1).strip()

            if name_match

            else "Patient name not found"

        )

    )


    id_match = _GROUP_MARKER.search(
        whole
    )


    record_id = (

        structured["patient_id"]

        or (

            id_match.group(1)

            if id_match

            else uuid.uuid4().hex[:12]

        )

    )


    patient = _create_patient(

        name,

        record_id,

        source

    )


    # Extract explicit KEY: VALUE fields.

    for match in _SINGLE_FIELD.finditer(whole):

        field_name = match.group(0).split(
            ":",
            1
        )[0]


        field = ALIAS_MAP.get(

            _field_key(field_name),

            _field_key(field_name)

        )


        if field != "name":

            patient["fields"][field] = normalize_text(

                match.group(1),

                500

            )


    # Add structured table fields.

    for key, value in structured["fields"].items():

        patient["fields"].setdefault(

            key,

            value

        )


    patient["field_sources"] = structured[
        "field_sources"
    ]


    # Create retrieval chunks.

    for text, location in parts:

        fragments = re.findall(

            r".{1,850}(?:\s|$)",

            text,

            re.S

        )


        for index, fragment in enumerate(fragments):

            _append_chunk(

                patient,

                fragment,

                source,

                f"{location}, part {index + 1}"

            )


    patient["chunks"] = patient["chunks"][:100]


    return [patient]


# ==========================================
# MAIN INGESTION ENTRY POINT
# ==========================================

def parse_file(
    contents: bytes,
    filename: str,
    mode: str
) -> list[dict]:

    suffix = Path(
        filename
    ).suffix.lower()


    if suffix == ".csv":

        return _parse_csv(

            contents,

            filename,

            mode

        )


    if suffix == ".xlsx":

        return _parse_xlsx(

            contents,

            filename,

            mode

        )


    if suffix == ".pdf":

        return _parse_text(

            _extract_pdf(contents),

            filename,

            mode

        )


    if suffix == ".docx":

        return _parse_text(

            _extract_docx(contents),

            filename,

            mode

        )


    raise HTTPException(
        422,
        "Unsupported file format."
    )